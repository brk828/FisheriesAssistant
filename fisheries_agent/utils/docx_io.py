"""Reads and writes Word (.docx) files as markdown-like text using python-docx (no Pandoc or shell needed)."""
import re
from pathlib import Path

INSTALL_HINT = "python-docx is not installed. Run verify_and_install_deps.py or: pip install python-docx"


def _docx():
    try:
        import docx
    except ImportError as e:
        raise RuntimeError(INSTALL_HINT) from e
    return docx


def read_docx(path: Path) -> str:
    """Body text in document order: headings as '#', list items as '-', tables as pipe rows."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = _docx().Document(str(path))
    out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            if m:
                out.append("#" * int(m.group(1)) + " " + text)
            elif style == "Title":
                out.append("# " + text)
            elif "List" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                out.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
    return "\n\n".join(out)


def write_docx(path: Path, content: str) -> None:
    """Markdown-lite to Word: '#' headings, '-'/'*' bullets, '1.' numbered items, pipe tables, **bold** and *italic*."""
    doc = _docx().Document()
    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    table.cell(r, c).text = text
            continue
        m = re.match(r"(#{1,6})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"\s*[-*]\s+", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"\s*\d+[.)]\s+", line):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        else:
            _runs(doc.add_paragraph(), line)
        i += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)
